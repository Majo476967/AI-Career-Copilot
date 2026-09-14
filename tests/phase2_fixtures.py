"""Synthetic documents and fake model outputs only."""
import copy
from io import BytesIO
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject

RESUME = "毕业于虚构学院软件工程专业。\n在虚构实验室实习，使用 SQL 完成订单查询练习。\n创建天气统计 Demo。\n了解 Python 基础。"
PROFILE = {
    "education": ["毕业于虚构学院软件工程专业。"],
    "internships": ["在虚构实验室实习，使用 SQL 完成订单查询练习。"],
    "projects": ["创建天气统计 Demo。"], "skills": ["SQL", "Python"],
    "capabilities": [{"name": "SQL能力", "level": 2, "evidence": [{
        "evidence_type": "internship", "content": "在虚构实验室实习，使用 SQL 完成订单查询练习。", "source": "resume"}]},
        {"name": "Python基础", "level": 1, "evidence": [{"evidence_type": "skill",
         "content": "了解 Python 基础。", "source": "resume"}]}],
}


def profile():
    return copy.deepcopy(PROFILE)


def jd_input(index=1):
    return f"虚构公司{index}招聘数据实习生。要求 SQL能力，能完成数据库查询练习；数据能力重要。"


def jd_output(index=1):
    return {"company": f"虚构公司{index}", "job_title": "数据实习生", "capabilities": [
        {"name": "SQL能力", "category": "technical", "importance": "must_have", "required_level": 2,
         "evidence": "要求 SQL能力，能完成数据库查询练习"},
        {"name": "数据能力", "category": "general", "importance": "important", "required_level": 1,
         "evidence": "数据能力重要"}]}


def docx_bytes(text=RESUME, *, table=False):
    document = Document()
    if table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = text
    else:
        document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(text="Synthetic resume SQL practice and project evidence", *, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(("BT /F1 12 Tf 40 700 Td (" + text + ") Tj ET").encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("test-password")
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()
